import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_proposal_docx(filename):
    doc = Document()
    
    # Page setup - 0.8 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles setup
    NAVY = RGBColor(30, 41, 59)      # #1E293B
    SLATE = RGBColor(71, 85, 105)    # #475569
    DARK = RGBColor(15, 23, 42)      # #0F172A
    ACCENT = RGBColor(16, 185, 129)  # Green
    
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = DARK
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)
    
    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("SMART SCHOOL FOOD OPERATIONS PLATFORM")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Commercial & Implementation Proposal | Brainwaves International School")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = SLATE
    r_sub.font.bold = True

    # Metadata Card (Table)
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Prepared For", "Board of Directors\nBrainwaves International School, Hapur, UP"),
         ("Prepared By", "Aboobacker Sidhiqe\nIndependent Product & Software Professional")],
        [("Date", "19 September 2026"),
         ("Document Version", "1.0 (Proposal & Commercial Framework)")]
    ]
    for row_idx, row in enumerate(meta_data):
        for col_idx, (label, val) in enumerate(row):
            cell = meta_table.cell(row_idx, col_idx)
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r_lbl = p.add_run(f"{label}: ")
            r_lbl.bold = True
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.color.rgb = SLATE
            r_v = p.add_run(val)
            r_v.font.size = Pt(9.5)
            r_v.font.color.rgb = DARK
    set_table_borders(meta_table, color="E2E8F0", sz="4")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def add_h1(text):
        p = doc.add_heading(level=1)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_heading(level=2)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = SLATE
        return p

    def format_table(table, col_widths, align_cols=None):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="E2E8F0", sz="4")
        for row_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if row_idx == 0:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for col_idx, cell in enumerate(row.cells):
                cell.width = Inches(col_widths[col_idx])
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                if align_cols and col_idx in align_cols:
                    p.alignment = align_cols[col_idx]
                if row_idx == 0:
                    set_cell_background(cell, "1E293B")
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(9.5)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "FFFFFF")
                    else:
                        set_cell_background(cell, "F8FAFC")
                    for r in p.runs:
                        r.font.size = Pt(9.5)

    # 1. Executive Summary
    add_h1("1. Executive Summary")
    doc.add_paragraph("Campus food services in modern educational institutions require balanced management of parent expectations, student well-being, kitchen efficiency, and financial control. The objective of the Smart School Food Operations Platform is not merely to digitize the canteen menu, but to introduce predictable demand forecasting before kitchen preparation begins, while offering parents an intuitive self-service portal to schedule, customize, and monitor their children's meals.")
    doc.add_paragraph("The platform connects parent pre-orders, kitchen production, classroom packaging, and administrative oversight into a single operational loop. Over time, the accumulated operational data provides school leadership with verifiable insights into food consumption trends, portion planning, waste elimination, and financial reporting.")
    
    add_h2("The Operational Shift")
    t1 = doc.add_table(rows=7, cols=3)
    t1_data = [
        ["Operational Dimension", "Traditional Manual Approach", "Smart School Food Platform"],
        ["Demand Forecasting", "Estimated intuitively; prone to over- or under-cooking", "Real-time aggregate order tallies locked prior to kitchen prep"],
        ["Batch Cooking", "Prepared by routine habit without confirmed counts", "Batches planned precisely around booked meals and breaks"],
        ["Order Coordination", "Paper slips, verbal requests, and manual registers", "Digital Kitchen Display System (KDS) with live order queues"],
        ["Parent Visibility", "No real-time tracking; requires admin intervention", "Transparent self-service: advance booking, repeat meals, status"],
        ["Fulfilment Audit", "Unrecorded handovers; difficult to trace missed meals", "Complete digital trail: Pre-order ➔ Batch Tally ➔ Portion & Pack ➔ Dispatched"],
        ["Operational Intelligence", "Fragmented or nonexistent historical records", "Structured data for menu popularity, peak slots, and waste reduction"]
    ]
    for r_i, row in enumerate(t1_data):
        for c_i, val in enumerate(row):
            t1.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t1, [1.5, 2.6, 2.8])

    # 2. Strategic Value
    add_h1("2. Strategic Value for School Leadership")
    points = [
        ("Predictive Demand & Waste Reduction: ", "Kitchen staff receive accurate pre-order tallies prior to the preparation cutoff. This eliminates speculative cooking and curbs avoidable raw material waste."),
        ("Institutional Allergen Safety & Liability Protection: ", "Automated cross-referencing of student medical profiles (Dairy, Gluten, Peanuts, Eggs) warns parents at checkout and stamps high-visibility red alerts on physical kitchen thermal stickers, mitigating school liability."),
        ("Cashless Campus & Financial Governance: ", "Eliminates physical cash handling by young students—removing lost pocket money, bullying risks, and cash hygiene concerns in the dining line—backed by an auditable digital transaction ledger."),
        ("Frictionless Multi-Sibling Family Hub: ", "Parents managing children in different grades order meals and select dietary preferences within a single unified checkout, maximizing institutional adoption."),
        ("Ground-Ready Usability for Local Canteen Staff: ", "Engineered with visual-first color progression, large token numbers, and 1-tap bilingual Hindi/English capability tailored specifically for ground canteen workers in Hapur."),
        ("Data Foundation for Procurement: ", "Historical consumption trends provide empirical data for seasonal purchasing and vendor negotiation.")
    ]
    for bold_pre, rest in points:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(bold_pre)
        r1.bold = True
        r1.font.color.rgb = NAVY
        p.add_run(rest)

    # 3. Proposed Operating Workflow
    add_h1("3. Proposed Operating Workflow & Lifecycle")
    doc.add_paragraph("The platform is engineered around an 8-stage operational cycle that links parents, administrative teams, and kitchen staff:")
    
    stages = [
        ["Stage", "Operational Action & Value"],
        ["1. Parent Ordering", "Parent selects student, meal item, date, and break slot via mobile portal."],
        ["2. Order Recording", "Order is validated against cutoff rules and logged in the central ledger."],
        ["3. Demand Aggregation", "Live kitchen summary calculates aggregate quantities required per meal & slot."],
        ["4. Kitchen Batch Cooking", "Canteen team cooks bulk quantities guided by KDS aggregate counters."],
        ["5. Portioning & Assembly", "Staff dish out hot portions into individual lunchboxes and apply thermal stickers."],
        ["6. Delivery / Pickup", "Packed crates are dispatched to classrooms or staged at collection counters."],
        ["7. Exception Logging", "Any cancellation, student absence, or shortage is flagged and audited."],
        ["8. Operational Analytics", "Data updates historical trend reports for future procurement and planning."]
    ]
    t_stages = doc.add_table(rows=len(stages), cols=2)
    for r_i, row in enumerate(stages):
        for c_i, val in enumerate(row):
            t_stages.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_stages, [1.8, 5.1])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_h2("Institutional Batch Cooking vs. Restaurant Ordering")
    p_batch_note = doc.add_paragraph()
    r_bn_b = p_batch_note.add_run("Important Operational Architecture: ")
    r_bn_b.bold = True
    r_bn_b.font.color.rgb = NAVY
    p_batch_note.add_run("Unlike commercial consumer food delivery apps (e.g., Swiggy or Zomato) that cook single dishes on demand as tickets arrive, school dining operates on predictive bulk catering. Canteen chefs prepare food in bulk batches before the recess bell based on locked morning pre-orders. The platform's individual order lifecycle is purposefully designed around batch verification, tray portioning, allergen labeling, and classroom crate sorting.")

    # 5-Stage Lifecycle
    add_h2("5-Stage Order Lifecycle")
    lifecycle_data = [
        ["Lifecycle Stage", "Parent Mobile App Status", "Kitchen / Operational Action"],
        ["1. NEW", "Order Confirmed", "Pre-order validated; added to morning batch tally"],
        ["2. ACCEPTED", "Preparing Fresh", "Canteen chef starts recipe bulk cooking for recess"],
        ["3. PREPARING", "Packing Lunch Box", "Hot portions dished into student's lunchbox"],
        ["4. PACKED", "Out for Delivery", "Lunchbox sealed with thermal sticker; sorted into class crate"],
        ["5. DELIVERED", "Delivered to Desk", "Crate runner places meal on child's desk for recess"]
    ]
    t_life = doc.add_table(rows=len(lifecycle_data), cols=3)
    for r_i, row in enumerate(lifecycle_data):
        for c_i, val in enumerate(row):
            t_life.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_life, [1.5, 2.5, 3.0])

    # 4. Exception Handling
    add_h1("4. Exception Handling & Operational Protocols")
    doc.add_paragraph("Real-world campus operations require clear protocols for holidays, student sickness, and stock variations. The platform handles these scenarios deterministically:")
    
    exceptions = [
        ["Operational Scenario", "Platform Handling", "Institutional Policy Application"],
        ["Cancellation Before Cutoff", "Permitted directly via Parent Portal.", "Credit/refund applied per school policy automatically."],
        ["Cancellation After Cutoff", "System locks changes; food prep has commenced.", "Exceptions require admin review; avoids kitchen loss."],
        ["Student Absence", "Parent cancels before morning cutoff window.", "Seamless advance credit; future attendance integration available."],
        ["Campus Closure / Emergency", "Admin closes service date in one click; bulk stops orders.", "Affected orders flagged for automated bulk credit/rescheduling."],
        ["Advance / Monthly Plan", "Parent books recurring calendar dates.", "Individual dates remain auditable; closures affect only specific days."],
        ["Kitchen Stockout / Shortage", "Item flagged as unfulfillable by kitchen manager.", "Prompt notification to parent; routed for substitution or credit."],
        ["Fulfilment Dispute", "Timestamped audit log inspected in Admin Console.", "School investigates exact handover stage with full accountability."],
        ["Duplicate Order Attempt", "Built-in duplicate detection warns parent.", "Prevents double-charging for same student on same date/slot."]
    ]
    t_ex = doc.add_table(rows=len(exceptions), cols=3)
    for r_i, row in enumerate(exceptions):
        for c_i, val in enumerate(row):
            t_ex.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_ex, [1.6, 2.6, 2.7])

    # 5. Platform Capabilities
    add_h1("5. Platform Capabilities & Architecture")
    caps = [
        ["Module / Capability", "Core Operational Functionality"],
        ["Parent Web App", "Mobile-first, progressive web portal for meal booking, order tracking, and balance inspection."],
        ["Multi-Child Hub", "Enables parents with multiple siblings to order across different grades in a single checkout."],
        ["Kitchen Display System (KDS)", "Large-format, real-time screen queue for kitchen staff showing live batch counts and order statuses."],
        ["Administration Console", "Complete portal for school managers to upload rosters, set menus, configure cutoff timers, and pull reports."],
        ["Classroom Delivery Workflow", "Groups completed orders by building, floor, and classroom section for organized delivery distribution."],
        ["Thermal Labeling Integration", "Instant printing of student-specific meal stickers (Student Name, Grade, Dietary tags, Meal Type)."],
        ["Dietary & Allergen Safeguards", "Flags student-specific restrictions (e.g., Jain, Nut Allergy, Gluten-free) directly on prep and packing tickets."],
        ["Audit & Analytics Foundation", "Generates exportable daily preparation summaries, sales reports, and exception logs."]
    ]
    t_cap = doc.add_table(rows=len(caps), cols=2)
    for r_i, row in enumerate(caps):
        for c_i, val in enumerate(row):
            t_cap.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_cap, [2.0, 4.9])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_h2("Institutional Safeguards & Operational Pillars")
    safeguards = [
        ["Strategic Pillar", "Institutional Value & Protection for the School"],
        ["Active Medical Allergen Shield", "Automatic conflict detection (Dairy, Peanuts, Gluten) warns parents at booking and prints prominent red alerts on kitchen meal labels to prevent life-threatening allergic reactions."],
        ["Cashless Campus Dining", "Prepaid Student Lunch Wallet and direct online checkout eliminate cash loss, theft, and physical currency handling in hygiene-sensitive dining spaces."],
        ["Unified Multi-Sibling Hub", "Parents manage multiple children across different grades and breaks in a single transaction with 1-tap meal copying, maximizing parent satisfaction and adoption."],
        ["Ground-Ready Visual KDS", "Zero-training interface featuring color-coded progression, large bold token numbers, and 1-tap bilingual Hindi/English capability tailored for local canteen workers in Hapur."]
    ]
    t_safe = doc.add_table(rows=len(safeguards), cols=2)
    for r_i, row in enumerate(safeguards):
        for c_i, val in enumerate(row):
            t_safe.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_safe, [2.2, 4.7])

    # Page Break before Commercials
    doc.add_page_break()

    # 6. Commercial Options
    add_h1("6. Commercial Options")
    doc.add_paragraph("To keep procurement straightforward, two software-and-service packages are available. Both options leverage the core cloud architecture; Option 2 introduces advanced classroom fulfilment, allergy safeguards, and multi-child workflows.")

    comms = [
        ["Feature / Deliverable", "Option 1: Essential Food App", "Option 2: Smart Campus Platform (Recommended)"],
        ["First-Year Total Investment", "₹59,000", "₹1,09,000"],
        ["Dedicated Parent Mobile Portal", "Included", "Included"],
        ["Advance Ordering & Cutoff Rules", "Included", "Included"],
        ["Kitchen Display System (Live Queue)", "Included", "Included"],
        ["Student Roster & Section Ingestion", "Included", "Included"],
        ["Standard Daily Order & Financial Reports", "Included", "Included"],
        ["Full Implementation & Staff Training", "Included", "Included"],
        ["Multi-Child / Sibling Unified Ordering", "—", "Included"],
        ["Allergen & Dietary Safety Tagging", "—", "Included"],
        ["Classroom Delivery & Logistics Module", "—", "Included"],
        ["Thermal Label / Sticker Printing Module", "—", "Included"],
        ["Advanced Demand & Operational Analytics", "—", "Included"],
        ["Hardware Requirements", "Existing Devices (BYOD)", "Existing Devices (BYOD) + Optional Printer"]
    ]
    t_comms = doc.add_table(rows=len(comms), cols=3)
    for r_i, row in enumerate(comms):
        for c_i, val in enumerate(row):
            t_comms.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_comms, [2.5, 2.2, 2.2], align_cols={1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.CENTER})

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    p_rat = doc.add_paragraph()
    r_rb = p_rat.add_run("Investment Rationale: ")
    r_rb.bold = True
    r_rb.font.color.rgb = NAVY
    p_rat.add_run("The pricing reflects a ready-to-deploy enterprise platform, inclusive of school-specific customization, secure cloud setup, testing, staff onboarding, and year-round support. Option 1 provides an agile entry point for core digital ordering, while Option 2 delivers a campus-wide food logistics system with allergy safety and multi-child convenience.")

    # 7. Annual Maintenance Contract (AMC)
    add_h1("7. Annual Maintenance Contract (AMC)")
    doc.add_paragraph("Following the completion of Year 1, ongoing platform availability, cloud hosting, and technical assistance are sustained via an annual maintenance contract pegged at 20% of the Year 1 investment:")

    amc_data = [
        ["Package", "Year 1 Investment", "Year 2+ Annual AMC (20%)"],
        ["Option 1 — Essential Food App", "₹59,000", "₹11,800 / year"],
        ["Option 2 — Smart Campus Platform", "₹1,09,000", "₹21,800 / year"]
    ]
    t_amc = doc.add_table(rows=len(amc_data), cols=3)
    for r_i, row in enumerate(amc_data):
        for c_i, val in enumerate(row):
            t_amc.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_amc, [2.5, 2.2, 2.2], align_cols={1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.CENTER})

    p_amc_scope = doc.add_paragraph()
    p_amc_scope.paragraph_format.space_before = Pt(4)
    r_as = p_amc_scope.add_run("AMC Scope of Service: ")
    r_as.bold = True
    r_as.font.color.rgb = NAVY
    p_amc_scope.add_run("Includes routine cloud maintenance, security updates, annual academic roster rollover (grade promotion), configuration tweaks, and SLA-backed technical assistance. Major new custom-engineered software modules or physical hardware replacements are quoted separately if requested.")

    # 8. Implementation Roadmap
    add_h1("8. Implementation Roadmap (7–14 Working Days)")
    doc.add_paragraph("Deployment is rapid and structured around school operational milestones once source information is provided:")

    roadmap = [
        ["Phase", "Timeline", "Core Activities & Deliverables"],
        ["1. Discovery & Setup", "Days 1–3", "Confirm meal cutoffs, service slots, menu pricing, and administrative contacts."],
        ["2. Ingestion & Branding", "Days 4–6", "Upload student master roster, grade sections, school logos, and menu items."],
        ["3. Testing & Walkthrough", "Days 7–9", "Verify kitchen queue, order cutoffs, cancellation rules, and mock orders."],
        ["4. Staff Training", "Days 10–11", "Remote walkthrough sessions for canteen operators and school administrative staff."],
        ["5. Pilot & Go-Live", "Days 12–14", "Phased rollout to selected classes or full campus launch with active launch support."]
    ]
    t_rm = doc.add_table(rows=len(roadmap), cols=3)
    for r_i, row in enumerate(roadmap):
        for c_i, val in enumerate(row):
            t_rm.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_rm, [1.6, 1.3, 4.0])

    # 9. Success Metrics
    add_h1("9. Success Metrics & Value Verification")
    doc.add_paragraph("Rather than making unsupported financial guarantees prior to deployment, the platform establishes clear operational benchmarks:")
    kpis = [
        ("Digital Adoption Rate: ", "Percentage of eligible students actively using digital ordering."),
        ("Demand Forecasting Lead Time: ", "Percentage of daily meals confirmed before kitchen preparation starts."),
        ("Order Accuracy: ", "Frequency of missing, mistaken, or disputed meals per service period."),
        ("Kitchen Efficiency: ", "Average turnaround time from order acceptance to packed status."),
        ("Avoidable Waste Tracking: ", "Measured variance between prepared, ordered, and unconsumed quantities.")
    ]
    for b_pre, r_txt in kpis:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(b_pre)
        r.bold = True
        r.font.color.rgb = NAVY
        p.add_run(r_txt)

    # 10. Governance & Responsibilities
    add_h1("10. Governance, Data Privacy & Responsibility Boundaries")
    gov_data = [
        ["Operational Domain", "Software Provider (Aboobacker Sidhiqe)", "School / Canteen Operator"],
        ["Software Uptime & Hosting", "Primary Responsible", "—"],
        ["Bug Fixes & System Maintenance", "Primary Responsible", "—"],
        ["Food Preparation, Hygiene & Quality", "—", "Solely Responsible"],
        ["Allergen Data Accuracy in Kitchen", "—", "Solely Responsible"],
        ["Physical Packing & Classroom Handover", "—", "Solely Responsible"],
        ["Refunds, Credits & Financial Settlement", "System Record Logging", "Policy Determination & Payout"],
        ["Parent Communication on Food Matters", "Technical Support Desk", "Canteen/Dining Policies"]
    ]
    t_gov = doc.add_table(rows=len(gov_data), cols=3)
    for r_i, row in enumerate(gov_data):
        for c_i, val in enumerate(row):
            t_gov.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_gov, [2.5, 2.2, 2.2], align_cols={1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.CENTER})

    p_gov_notes = doc.add_paragraph()
    p_gov_notes.paragraph_format.space_before = Pt(4)
    p_gov_notes.add_run("Data Privacy & Ownership: ").bold = True
    p_gov_notes.add_run("All student and institutional transaction data remains the exclusive property of Brainwaves International School. Data is stored securely and never shared with third parties. Food revenues are collected directly by the school's designated accounts; the software provider does not hold food funds.")

    # 11. Payment Terms
    add_h1("11. Payment Terms & Milestones")
    pay_data = [
        ["Milestone", "Share (%)", "Option 1 (₹)", "Option 2 (₹)", "Trigger Condition"],
        ["1. Project Inception", "50%", "₹29,500", "₹54,500", "Agreement execution and onboarding kickoff"],
        ["2. Configuration & Staging", "30%", "₹17,700", "₹32,700", "Roster/menu configured; staging testing approved"],
        ["3. Go-Live & Handover", "20%", "₹11,800", "₹21,800", "Staff training completed and campus launch"]
    ]
    t_pay = doc.add_table(rows=len(pay_data), cols=5)
    for r_i, row in enumerate(pay_data):
        for c_i, val in enumerate(row):
            t_pay.cell(r_i, c_i).paragraphs[0].add_run(val)
    format_table(t_pay, [1.5, 0.8, 1.1, 1.1, 2.4], align_cols={1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.RIGHT, 3: WD_ALIGN_PARAGRAPH.RIGHT})

    # 12. Next Steps & Sign-off
    add_h1("12. Next Steps to Proceed")
    steps = [
        "1. Confirm commercial selection between Option 1 (Essential) and Option 2 (Smart Campus).",
        "2. Confirm institutional policies (meal cutoffs, cancellation windows, credit rules).",
        "3. Designate a school coordinator for roster and menu configuration.",
        "4. Sign formal work order and initiate Phase 1 deployment."
    ]
    for step in steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(step)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # Sign-off box
    sign_table = doc.add_table(rows=1, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sign_table, color="CBD5E1", sz="4")
    
    c1 = sign_table.cell(0, 0)
    set_cell_background(c1, "F8FAFC")
    set_cell_margins(c1, top=120, bottom=120, left=150, right=150)
    p1 = c1.paragraphs[0]
    p1.add_run("Submitted by:\n").bold = True
    p1.add_run("Aboobacker Sidhiqe\nIndependent Freelance Product & Software Professional\nDate: 19 September 2026")
    
    c2 = sign_table.cell(0, 1)
    set_cell_background(c2, "F8FAFC")
    set_cell_margins(c2, top=120, bottom=120, left=150, right=150)
    p2 = c2.paragraphs[0]
    p2.add_run("Accepted on behalf of Brainwaves Int. School:\n\n").bold = True
    p2.add_run("Name & Title: _________________________________\n\nSignature: __________________  Date: ____________")

    sign_table.rows[0].cells[0].width = Inches(3.4)
    sign_table.rows[0].cells[1].width = Inches(3.5)

    doc.save(filename)
    print(f"Successfully generated {filename}")

if __name__ == "__main__":
    out_path = r"e:\Project\school-food-app\Smart_School_Food_Operations_Proposal.docx"
    build_proposal_docx(out_path)
